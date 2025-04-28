import streamlit as st
import gspread
from oauth2client.service_account import ServiceAccountCredentials
from utils.genai_resume import generate_resume
from utils.ml_model import match_resume_with_job
from utils.generate_pdf import generate_pdf
from utils.resume_templates import get_resume_template
from PyPDF2 import PdfReader
from docx import Document

# ---------------- Visitor Counter Setup ----------------

# Authenticate with Google Sheets
scope = ["https://spreadsheets.google.com/feeds", "https://www.googleapis.com/auth/drive"]
creds = ServiceAccountCredentials.from_json_keyfile_name('credentials.json', scope)
client = gspread.authorize(creds)

# Open your Google Sheet
sheet = client.open('VisitorCountsForAutoResume').sheet1  # Change if your sheet name is different

# Read current visitor count
count = int(sheet.cell(1, 1).value)

# Update the count
count += 1
sheet.update_cell(1, 1, count)

# Show visitor count in sidebar
st.sidebar.title("👥 Visitors")
st.sidebar.write(f"Total Visitors: {count}")

# ---------------- Streamlit App Setup ----------------

# Set the page title
st.set_page_config(page_title="Auto Resume Creator")

# Header section
st.title("Auto Resume Creator")
st.markdown("Create your resume based on a job description and apply for jobs effortlessly!")

# Function to extract text from PDF
def extract_pdf_text(file):
    pdf_reader = PdfReader(file)
    text = ""
    for page in pdf_reader.pages:
        text += page.extract_text()
    return text

# Function to extract text from DOCX
def extract_docx_text(file):
    doc = Document(file)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text

# Upload job description
uploaded_file = st.file_uploader("Upload Job Description (Word or PDF)", type=["pdf", "docx"])
job_desc_text = ""

if uploaded_file is not None:
    if uploaded_file.type == "application/pdf":
        job_desc_text = extract_pdf_text(uploaded_file)
    elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        job_desc_text = extract_docx_text(uploaded_file)

# Show job description
if job_desc_text:
    st.subheader("Job Description")
    st.write(job_desc_text)

# Generate Resume Button
if st.button("Generate Resume"):
    if job_desc_text:
        st.subheader("Resume Generation")
        resume = generate_resume(job_desc_text)
        st.text_area("Generated Resume", value=resume, height=300)
        if st.button("Download Resume (PDF)"):
            generate_pdf(resume)  # Make sure generate_pdf is properly defined

# Match resume button (If the user has a resume)
uploaded_resume = st.file_uploader("Upload Your Resume (Word or PDF)", type=["pdf", "docx"])

if uploaded_resume is not None:
    resume_text = ""
    if uploaded_resume.type == "application/pdf":
        resume_text = extract_pdf_text(uploaded_resume)
    elif uploaded_resume.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        resume_text = extract_docx_text(uploaded_resume)
    
    # Job fit matching
    if st.button("Match Resume with Job Description"):
        if job_desc_text:  # Check if job description is available
            match_score = match_resume_with_job(resume_text, job_desc_text)
            st.write(f"Job Fit Score: {match_score * 100:.2f}%")
        else:
            st.error("Please upload a job description before matching.")
